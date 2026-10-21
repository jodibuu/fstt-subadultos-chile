"""Orientación anatómica de mallas en formato PLY.

Aplica X +90° y posteriormente Y −90° sobre vértices y normales. El proceso
debe realizarse una sola vez, preferentemente sobre copias respaldadas.
"""
import gc
import os
import sys
import shutil
import argparse
import tempfile
import time
import numpy as np
import openpyxl  # pip install openpyxl

# ─── Configuración por defecto ───────────────────────────────
from config_pipeline import RESULTADOS_FOLDER
DEFAULT_EXCEL = os.environ.get("FSTT_EXCEL_ADDON", str(RESULTADOS_FOLDER / "excel_mediciones_local.xlsx"))

# Nombres de columnas (deben coincidir con el Excel del proyecto)
COL_INCLUIR = "¿Incluir?"
COL_MEJOR = "¿Es la mejor carpeta del paciente? (1 = Sí)"
COL_RUTA_BLANDO = "Ruta malla blando"
COL_RUTA_HUESO = "Ruta malla hueso"
COL_PACIENTE = "Paciente"
COL_OJO = "Ojo"

SHEET_CANDIDATES = ["Estadísticas", "Casos", "Sheet1"]


# ─── Matriz de rotación combinada ────────────────────────────
# Primero X +90°, luego Y -90° (mismo orden que el addon)
#
#   Rx(90)  = [[1,  0,  0],       Ry(-90) = [[ 0, 0, -1],
#              [0,  0, -1],                   [ 0, 1,  0],
#              [0,  1,  0]]                   [ 1, 0,  0]]
#
#   M = Ry(-90) @ Rx(90) = [[ 0, -1,  0],
#                            [ 0,  0, -1],
#                            [ 1,  0,  0]]
#

def _build_rotation_matrix():
    """Construye la matriz 3×3 de rotación combinada X+90° luego Y-90°."""
    cx, sx = np.cos(np.radians(90)), np.sin(np.radians(90))
    Rx = np.array([
        [1,   0,    0],
        [0,  cx,  -sx],
        [0,  sx,   cx],
    ])

    cy, sy = np.cos(np.radians(-90)), np.sin(np.radians(-90))
    Ry = np.array([
        [ cy,  0,  sy],
        [  0,  1,   0],
        [-sy,  0,  cy],
    ])

    return Ry @ Rx  # primero Rx, luego Ry


ROT_MATRIX = _build_rotation_matrix()


# ─── Lectura del Excel ───────────────────────────────────────

def _is_one(v):
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    try:
        return float(v) == 1.0
    except Exception:
        return str(v).strip().lower() in {"1", "si", "sí", "true", "yes"}


def leer_rutas_desde_excel(excel_path: str, pacientes: set[str] | None = None) -> list[dict]:
    """Lee un Excel QC/maestro y retorna rutas de mallas de las filas Mejor=1.

    Acepta hojas ``Estadísticas``, ``Casos`` o ``Sheet1``. Si existe la columna
    ``¿Incluir?`` también exige Incluir=1. ``pacientes`` permite limitar el
    procesamiento a IDs concretos.
    """
    if not os.path.exists(excel_path):
        print(f"ERROR: Excel no encontrado: {excel_path}")
        sys.exit(1)

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    sheet_name = next((n for n in SHEET_CANDIDATES if n in wb.sheetnames), None)
    if sheet_name is None:
        print(f"ERROR: No se encontró una hoja compatible. Hojas: {wb.sheetnames}")
        wb.close()
        sys.exit(1)

    ws = wb[sheet_name]
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = {v: i for i, v in enumerate(header_row) if v is not None}

    required = [COL_MEJOR, COL_RUTA_BLANDO, COL_RUTA_HUESO, COL_PACIENTE]
    missing = [c for c in required if c not in headers]
    if missing:
        print(f"ERROR: Faltan columnas requeridas: {missing}")
        wb.close()
        sys.exit(1)

    i_mej = headers[COL_MEJOR]
    i_rb = headers[COL_RUTA_BLANDO]
    i_rh = headers[COL_RUTA_HUESO]
    i_pac = headers[COL_PACIENTE]
    i_ojo = headers.get(COL_OJO)
    i_inc = headers.get(COL_INCLUIR)

    casos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        paciente = str(row[i_pac] or "").strip()
        if not paciente:
            continue
        if pacientes and paciente not in pacientes:
            continue
        if not _is_one(row[i_mej]):
            continue
        if i_inc is not None and not _is_one(row[i_inc]):
            continue
        casos.append({
            "paciente": paciente,
            "ojo": str(row[i_ojo] or "") if i_ojo is not None else "",
            "ruta_hueso": str(row[i_rh] or "").strip(),
            "ruta_blando": str(row[i_rb] or "").strip(),
        })

    wb.close()
    print(f"Hoja usada: {sheet_name}")
    return casos


# ─── Parser rápido de PLY binario (sin plyfile) ─────────────

_PLY_TYPE_MAP = {
    'char': 'i1', 'uchar': 'u1', 'int8': 'i1', 'uint8': 'u1',
    'short': 'i2', 'ushort': 'u2', 'int16': 'i2', 'uint16': 'u2',
    'int': 'i4', 'uint': 'u4', 'int32': 'i4', 'uint32': 'u4',
    'float': 'f4', 'float32': 'f4', 'double': 'f8', 'float64': 'f8',
}


def _parse_ply_header(filepath: str):
    """Lee solo el encabezado de un PLY binario.

    Returns:
        (header_size, byte_order, n_verts, vertex_dtype, has_normals)
    """
    with open(filepath, 'rb') as f:
        header_lines: list[bytes] = []
        while True:
            line = f.readline()
            if not line:
                raise ValueError("Fin de archivo antes de end_header")
            header_lines.append(line)
            if line.strip() == b'end_header':
                break
        header_size = f.tell()

    text = b''.join(header_lines).decode('ascii', errors='replace')
    lines = text.splitlines()

    # Formato binario
    bo = None
    for l in lines:
        if l.startswith('format '):
            fmt = l.split()[1]
            if fmt == 'binary_little_endian':
                bo = '<'
            elif fmt == 'binary_big_endian':
                bo = '>'
            else:
                raise ValueError(f"Formato PLY no soportado para modo rápido: {fmt}")
            break
    if bo is None:
        raise ValueError("No se encontró línea 'format' en el encabezado PLY")

    # Parsear elementos y propiedades
    elements: list[dict] = []
    cur: dict | None = None
    for l in lines:
        if l.startswith('element '):
            p = l.split()
            cur = {'name': p[1], 'count': int(p[2]), 'props': [], 'has_list': False}
            elements.append(cur)
        elif l.startswith('property ') and cur is not None:
            p = l.split()
            if p[1] == 'list':
                cur['has_list'] = True
            else:
                cur['props'].append((p[1], p[2]))  # (tipo, nombre)

    # El elemento vertex debe ser el primero (estándar PLY)
    if not elements or elements[0]['name'] != 'vertex':
        raise ValueError("El elemento 'vertex' debe ser el primero en el PLY")

    ve = elements[0]
    if ve['has_list']:
        raise ValueError("El elemento vertex tiene propiedades list")

    # Construir numpy dtype
    dt_fields = []
    prop_names = []
    for ptype, pname in ve['props']:
        np_t = _PLY_TYPE_MAP.get(ptype)
        if np_t is None:
            raise ValueError(f"Tipo PLY desconocido: {ptype}")
        dt_fields.append((pname, bo + np_t))
        prop_names.append(pname)

    vertex_dtype = np.dtype(dt_fields)
    n_verts = ve['count']
    has_normals = all(n in prop_names for n in ('nx', 'ny', 'nz'))

    return header_size, bo, n_verts, vertex_dtype, has_normals


# ─── Rotación de PLY (I/O binario directo, sin plyfile) ──────

def rotar_ply(filepath: str, make_backup: bool = True, dry_run: bool = False) -> bool:
    """Aplica la rotación combinada X+90° Y-90° a un archivo PLY.

    Lee el encabezado, parsea vértices como numpy structured array,
    aplica la rotación y escribe todo de vuelta. ~5-10× más rápido
    que usar plyfile.

    Returns:
        True si se procesó, False si se saltó o hubo error.
    """
    if not os.path.exists(filepath):
        print(f"  SKIP — archivo no existe: {filepath}")
        return False

    if dry_run:
        print(f"  DRY-RUN — se rotaría: {filepath}")
        return True

    try:
        header_size, bo, n_verts, vdtype, has_normals = _parse_ply_header(filepath)
    except ValueError as e:
        print(f"  ERROR encabezado PLY: {e}")
        return False

    vertex_bytes_len = n_verts * vdtype.itemsize

    # Leer archivo en 3 bloques: header, vértices (raw), resto (caras, etc.)
    try:
        with open(filepath, 'rb') as f:
            header_raw = f.read(header_size)
            vertex_raw = f.read(vertex_bytes_len)
            rest_raw = f.read()
    except Exception as e:
        print(f"  ERROR leyendo archivo: {e}")
        return False

    if len(vertex_raw) != vertex_bytes_len:
        print(f"  ERROR — datos truncados ({len(vertex_raw)} vs {vertex_bytes_len} esperados)")
        return False

    # Parsear vértices como structured array
    verts = np.frombuffer(vertex_raw, dtype=vdtype).copy()
    del vertex_raw

    # Backup antes de sobreescribir
    if make_backup:
        backup = filepath + ".bak"
        if not os.path.exists(backup):
            shutil.copy2(filepath, backup)
            print(f"  Backup → {backup}")

    # Rotar xyz
    xyz = np.column_stack([
        verts['x'].astype(np.float64),
        verts['y'].astype(np.float64),
        verts['z'].astype(np.float64),
    ])
    rr = (ROT_MATRIX @ xyz.T).T
    verts['x'] = rr[:, 0].astype(verts['x'].dtype)
    verts['y'] = rr[:, 1].astype(verts['y'].dtype)
    verts['z'] = rr[:, 2].astype(verts['z'].dtype)
    del xyz, rr

    # Rotar normales si existen
    if has_normals:
        nrm = np.column_stack([
            verts['nx'].astype(np.float64),
            verts['ny'].astype(np.float64),
            verts['nz'].astype(np.float64),
        ])
        nr = (ROT_MATRIX @ nrm.T).T
        verts['nx'] = nr[:, 0].astype(verts['nx'].dtype)
        verts['ny'] = nr[:, 1].astype(verts['ny'].dtype)
        verts['nz'] = nr[:, 2].astype(verts['nz'].dtype)
        del nrm, nr

    # Guardar: header original + vértices rotados + resto (caras, etc.)
    try:
        dirn = os.path.dirname(filepath)
        fd, tmp_path = tempfile.mkstemp(suffix=".ply", dir=dirn)
        try:
            with os.fdopen(fd, 'wb') as f:
                f.write(header_raw)
                f.write(verts.tobytes())
                f.write(rest_raw)
            del verts, header_raw, rest_raw
            os.replace(tmp_path, filepath)
        except Exception:
            try:
                os.unlink(tmp_path)
            except OSError:
                pass
            raise
        print(f"  OK — {n_verts} vértices rotados" + (" + normales" if has_normals else ""))
        return True
    except Exception as e:
        print(f"  ERROR guardando PLY: {e}")
        return False


# ─── Main ────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Pre-rota mallas PLY del proyecto (X +90°, Y -90°) para que "
                    "el addon no necesite hacerlo al importar."
    )
    parser.add_argument(
        "--excel", default=DEFAULT_EXCEL,
        help=f"Ruta al archivo Excel del proyecto (default: {DEFAULT_EXCEL})"
    )
    parser.add_argument(
        "--dry-run", action="store_true",
        help="Solo muestra qué archivos se procesarían, sin modificar nada"
    )
    parser.add_argument(
        "--no-backup", action="store_true",
        help="No crear archivos .bak antes de sobreescribir"
    )
    parser.add_argument(
        "--resume", action="store_true",
        help="Compatibilidad: saltar archivos que ya tienen .bak"
    )
    parser.add_argument(
        "--pacientes", nargs="+", default=None,
        help="Procesar solo estos pacientes, por ejemplo: --pacientes CASO_001 CASO_002"
    )
    args = parser.parse_args()

    print("=" * 60)
    print("  ORIENTACIÓN ANATÓMICA DE MALLAS")
    print("  Rotación: X +90°, luego Y -90°")
    print(f"  Excel: {args.excel}")
    if args.dry_run:
        print("  MODO: dry-run (no se modificará nada)")
    if args.resume:
        print("  MODO: resume (se saltarán archivos con .bak existente)")
    print("=" * 60)

    pacientes = set(args.pacientes) if args.pacientes else None
    casos = leer_rutas_desde_excel(args.excel, pacientes=pacientes)
    filtro_txt = ", ".join(args.pacientes) if args.pacientes else "TODOS"
    print(f"  Pacientes: {filtro_txt}")
    print(f"\nCasos encontrados (Mejor=1 y, si existe, Incluir=1): {len(casos)}\n")

    if not casos:
        print("No hay casos para procesar.")
        return

    # Recolectar todas las rutas únicas (un mismo PLY podría aparecer en varios casos)
    rutas_unicas: dict[str, str] = {}  # ruta → "paciente/ojo (tipo)"
    for c in casos:
        for tipo, ruta in [("hueso", c["ruta_hueso"]), ("blando", c["ruta_blando"])]:
            if ruta and ruta not in rutas_unicas:
                rutas_unicas[ruta] = f"{c['paciente']}/{c['ojo']} ({tipo})"

    total = len(rutas_unicas)
    print(f"Archivos PLY únicos a procesar: {total}\n")

    ok = 0
    skip = 0
    err = 0
    t_start = time.time()

    for i, (ruta, desc) in enumerate(sorted(rutas_unicas.items()), 1):
        # SEGURIDAD: .bak significa que esta PLY ya fue pre-rotada antes.
        # Se omite siempre para evitar aplicar la rotación dos veces.
        if os.path.exists(ruta + ".bak"):
            print(f"[{i}/{total}] [{desc}] SKIP (ya tiene .bak; no se rota dos veces)")
            skip += 1
            continue

        t0 = time.time()
        print(f"[{i}/{total}] [{desc}] {os.path.basename(ruta)}")
        result = rotar_ply(ruta, make_backup=not args.no_backup, dry_run=args.dry_run)
        dt = time.time() - t0
        if result:
            ok += 1
            print(f"  ({dt:.1f}s)")
        elif not os.path.exists(ruta):
            skip += 1
        else:
            err += 1
        gc.collect()

    elapsed = time.time() - t_start
    print("\n" + "=" * 60)
    print(f"  Resultado: {ok} procesados, {skip} saltados, {err} errores")
    print(f"  Tiempo total: {elapsed:.0f}s ({elapsed/60:.1f} min)")
    if args.dry_run:
        print("  (dry-run — ningún archivo fue modificado)")
    if not args.no_backup and not args.dry_run:
        print("  Los originales se guardaron como .bak")
    print("=" * 60)


if __name__ == "__main__":
    main()
